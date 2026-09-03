"""
Markdown -> .docx and -> styled .html for the syllabus and lesson plans.

    python tools/build_docs.py

Reads syllabus/*.md and lessons/*.md, writes a .docx next to each source and an
.html copy into docs/ (served by GitHub Pages next to the decks).
"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCES = sorted((ROOT / 'syllabus').glob('*.md')) + sorted((ROOT / 'lessons').glob('*.md'))


# ───────────────────────── docx ─────────────────────────
def md_to_docx(md: str, out: Path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    for s in ('Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[s].font.color.rgb = RGBColor(0x00, 0x0B, 0x1C)
        doc.styles[s].font.name = 'Calibri'

    def add_runs(par, text):
        # **bold**, `code`
        for part in re.split(r'(\*\*.+?\*\*|`[^`]+`)', text):
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith('`') and part.endswith('`'):
                r = par.add_run(part[1:-1]); r.font.name = 'Consolas'
            else:
                par.add_run(part)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            doc.add_heading(line.lstrip('#').strip(), level=min(level, 3))
        elif line.startswith('>'):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            r = p.add_run(line.lstrip('> ').strip()); r.italic = True
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = 'Light Grid Accent 1' if 'Light Grid Accent 1' in [s.name for s in doc.styles] else 'Table Grid'
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    cell.text = ''
                    par = cell.paragraphs[0]
                    add_runs(par, row[ci] if ci < len(row) else '')
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
            doc.add_paragraph()
            continue
        elif re.match(r'^\s*[-*] ', line):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^\s*[-*] ', '', line))
        elif re.match(r'^\s*\d+\. ', line):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\s*\d+\. ', '', line))
        else:
            # merge soft-wrapped lines of one paragraph
            buf = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(#|>|\||\s*[-*] |\s*\d+\. )', lines[i + 1]):
                i += 1; buf.append(lines[i].rstrip())
            p = doc.add_paragraph()
            add_runs(p, ' '.join(buf))
        i += 1
    doc.save(out)


# ───────────────────────── html ─────────────────────────
HTML = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<link rel="stylesheet" href="vendor/ait4x-colors_and_type.css">
<style>
@font-face{{font-family:'Inter';src:url('vendor/fonts/Inter-Variable.ttf') format('truetype');font-weight:100 900}}
@font-face{{font-family:'JetBrains Mono';src:url('vendor/fonts/JetBrainsMono-Variable.ttf') format('truetype');font-weight:100 800}}
body{{max-width:980px;margin:0 auto;padding:48px 24px 96px;color:var(--fg-1)}}
.eyebrow{{margin-bottom:24px}} h1{{margin:0 0 24px;font-size:clamp(36px,6vw,72px)}} h2{{margin:56px 0 16px}} h3{{margin:32px 0 8px}}
p,li{{font-size:17px;line-height:1.55;color:var(--fg-2)}} li{{margin:4px 0}}
table{{border-collapse:collapse;width:100%;margin:16px 0 24px;font-size:15px}} th,td{{text-align:left;vertical-align:top;padding:10px 12px;border-top:1px solid var(--line)}} th{{font-family:var(--font-mono);font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:var(--fg-3)}}
blockquote{{margin:24px 0;padding:0 0 0 20px;border-left:4px solid var(--accent-warm);font-family:var(--font-display);font-weight:800;font-size:24px;line-height:1.15;letter-spacing:-.02em;color:var(--fg-1)}} blockquote p{{font-size:inherit;color:inherit;line-height:inherit;margin:0}}
code{{font-family:var(--font-mono);font-size:.92em;background:var(--bg-2);padding:1px 5px}}
a{{text-decoration:underline;text-underline-offset:.12em}} a:hover{{color:var(--accent-deep)}}
nav{{display:flex;gap:24px;align-items:baseline;margin-bottom:48px}} nav a{{font-family:var(--font-mono);font-size:12px;letter-spacing:.14em;text-transform:uppercase;text-decoration:none;color:var(--fg-3)}}
nav a.wordmark{{font-family:var(--font-display);font-size:28px;font-weight:900;letter-spacing:-.04em;text-transform:none;color:var(--fg-1)}}
</style></head><body>
<nav><a href="index.html" class="wordmark" style="text-decoration:none;color:var(--fg-1)">a<span class="dot"></span>t4<span class="x">x</span></a><a href="index.html">SD2112</a><a href="week01/">Week 1 slides</a><a href="syllabus.html">Syllabus</a><a href="week01-lesson-plan.html">Lesson plan</a></nav>
<div class="eyebrow">POLYU SCHOOL OF DESIGN · SD2112 · 2026/27</div>
{body}
</body></html>
"""


def md_to_html(md: str, out: Path, title: str):
    import markdown
    body = markdown.markdown(md, extensions=['tables'])
    out.write_text(HTML.format(title=title, body=body), encoding='utf-8')


if __name__ == '__main__':
    names = {'SD2112-syllabus-2026.md': ('syllabus.html', 'SD2112 · Syllabus 2026/27'),
             'week01-lesson-plan.md': ('week01-lesson-plan.html', 'SD2112 · Week 1 lesson plan')}
    for src in SOURCES:
        md = src.read_text(encoding='utf-8')
        docx_out = src.with_suffix('.docx')
        md_to_docx(md, docx_out)
        html_name, title = names.get(src.name, (src.stem + '.html', src.stem))
        md_to_html(md, ROOT / 'docs' / html_name, title)
        print(f'{src.name} -> {docx_out.name}, docs/{html_name}')
