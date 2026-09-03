"""
deckgen — one slide spec, three outputs.

A deck is a plain Python structure (see deck/week01.py). Every slide is
laid out once, on a 1920 x 1080 px canvas, into a list of primitive
elements (text, rect, line, image, figure, embed). Three backends then
render the same elements:

  html   — a reveal.js deck for GitHub Pages (docs/<deck>/index.html)
  pptx   — an editable PowerPoint via python-pptx, post-processed by
           tools/classpoint/build.py (ait4x master, animations, ClassPoint)
  png    — a PIL preview of every slide + a text-overflow check

Design tokens come from the ait4x design system (docs/vendor/ait4x-colors_and_type.css).
Font names follow tools/PPTX-EXPORT.md: the family name carries the weight.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FONT_DIR = ROOT / 'tools' / 'fonts'
W, H = 1920, 1080
PX = 6350  # EMU per px on a 1920x1080 (13.333 x 7.5 in) slide

# ───────────────────────── tokens (ait4x / PolyU Design) ─────────────────────────
INK, WHITE, PAPER, GRAY, TEAL = '#000B1C', '#FFFFFF', '#F4F4F2', '#BBBCB9', '#64C2C3'
TXT, MUTED, LINE, LINE_STRONG = '#2A323D', '#5C6470', '#E1E1DE', '#1F2832'
ORANGE, VIOLET, PINK, YELLOW, GREEN, BLUE, DEEP_TEAL, RED = '#ED6D24', '#943890', '#E94D7F', '#F6AD00', '#6FBA2C', '#146AB5', '#00544C', '#E42519'
MUTED_ON_INK, LIGHT_ON_INK = '#B3B7BE', '#D3E7E8'
# secondary matrix value scales used for bands
YELLOWS = ['#F6AD00', '#F0BD60', '#F4CE89', '#F8DEB1']
VIOLETS = ['#943890', '#A167A1', '#B68EBB', '#CDB5D4', '#E5DAEB', '#F2ECF5']
TEALS = ['#64C2C3', '#9FCED0', '#B9DBDC', '#D3E7E8', '#E9F3F4']
ORANGES = ['#ED6D24', '#E38E5D', '#EBAC83', '#F2C9AC', '#F9E5D6']
PINKS = ['#E94D7F', '#DF7F99', '#E7A3B4', '#EFC4CF', '#F7E3E8']

# font key -> (PowerPoint family, bold flag, css weight, PIL variation, css family, is_mono)
FONTS = {
    'black':    ('Inter Black', False, 900, 'Black', 'Inter', False),
    'xbold':    ('Inter ExtraBold', False, 800, 'ExtraBold', 'Inter', False),
    'bold':     ('Inter', True, 700, 'Bold', 'Inter', False),
    'semibold': ('Inter SemiBold', False, 600, 'SemiBold', 'Inter', False),
    'medium':   ('Inter Medium', False, 500, 'Medium', 'Inter', False),
    'body':     ('Inter', False, 400, 'Regular', 'Inter', False),
    'light':    ('Inter Light', False, 300, 'Light', 'Inter', False),
    'mono':     ('JetBrains Mono', False, 400, 'Regular', 'JetBrains Mono', True),
    'monomed':  ('JetBrains Mono Medium', False, 500, 'Medium', 'JetBrains Mono', True),
    'monobold': ('JetBrains Mono', True, 700, 'Bold', 'JetBrains Mono', True),
}


# ───────────────────────── element model ─────────────────────────
@dataclass
class Run:
    text: str
    font: str = 'body'
    size: int = 36
    color: str = INK
    spc: float = 0.0        # letter spacing in em
    caps: bool = False
    italic: bool = False
    url: str | None = None


@dataclass
class Para:
    runs: list
    align: str = 'l'        # l c r
    lh: float = 1.2         # line height, multiple of font size
    before: int = 0         # space before, px
    bullet: bool = False    # orange dot bullet (level 2 of the ait4x master)


@dataclass
class Text:
    x: int; y: int; w: int; h: int
    paras: list
    valign: str = 't'       # t m b
    name: str = ''
    kind: str = 'text'


@dataclass
class Rect:
    x: int; y: int; w: int; h: int
    fill: str | None = None
    stroke: str | None = None
    stroke_w: int = 2
    name: str = ''
    kind: str = 'rect'


@dataclass
class Image:
    x: int; y: int; w: int; h: int
    src: str                # path relative to deck/assets or absolute
    fit: str = 'cover'      # cover | contain
    name: str = ''
    kind: str = 'image'


@dataclass
class Figure:
    """A drawn illustration: svg for html, png for pptx/preview."""
    x: int; y: int; w: int; h: int
    svg: str
    png: str                # path to a rendered png
    name: str = ''
    kind: str = 'figure'


@dataclass
class Embed:
    """YouTube embed. html: iframe; pptx/png: thumbnail + link."""
    x: int; y: int; w: int; h: int
    yt: str
    thumb: str | None = None
    name: str = ''
    kind: str = 'embed'


@dataclass
class Slide:
    bg: str = WHITE
    els: list = field(default_factory=list)
    notes: str = ''
    cp: dict | None = None      # ClassPoint activity for build.py
    title: str = ''             # for the outline / html title
    html_only: list = field(default_factory=list)  # elements only for html (e.g. live demos)


# ───────────────────────── inline markup ─────────────────────────
# **bold**  ·  [text](url)  ·  {orange:text} {teal:text} {muted:text} {mono:text}
_TOKEN = re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|\{[a-z]+:[^}]+\})')
_COLORS = {'orange': ORANGE, 'teal': TEAL, 'muted': MUTED, 'violet': VIOLET, 'pink': PINK,
           'white': WHITE, 'ink': INK, 'green': GREEN, 'yellow': YELLOW, 'blue': BLUE}


def runs(text, font='body', size=36, color=INK, spc=0.0, caps=False, bold_font=None, italic=False):
    """Parse light inline markup into Run objects."""
    bold_font = bold_font or ('bold' if font in ('body', 'light', 'medium') else font)
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            out.append(Run(part[2:-2], bold_font, size, color, spc, caps, italic))
        elif part.startswith('[') and '](' in part:
            label, url = part[1:-1].split('](')
            out.append(Run(label, font, size, color, spc, caps, italic, url=url))
        elif part.startswith('{') and ':' in part:
            key, val = part[1:-1].split(':', 1)
            if key == 'mono':
                out.append(Run(val, 'mono', size, color, spc, caps, italic))
            elif key == 'bold':
                out.append(Run(val, bold_font, size, color, spc, caps, italic))
            else:
                out.append(Run(val, font, size, _COLORS.get(key, color), spc, caps, italic))
        else:
            out.append(Run(part, font, size, color, spc, caps, italic))
    return out


def T(x, y, w, h, text, font='body', size=36, color=INK, lh=1.2, align='l', valign='t', spc=0.0, caps=False, name='', italic=False):
    """Single-paragraph text box; '\n' starts a new paragraph."""
    paras = [Para(runs(line, font, size, color, spc, caps, italic=italic), align, lh) for line in text.split('\n')]
    return Text(x, y, w, h, paras, valign, name)


def P(text, font='body', size=36, color=INK, lh=1.4, align='l', before=0, bullet=False, spc=0.0, caps=False):
    return Para(runs(text, font, size, color, spc, caps), align, lh, before, bullet)


def eyebrow(x, y, text, color=MUTED, w=1500, size=24, name=''):
    return T(x, y, w, 40, text, 'monomed', size, color, lh=1.2, spc=0.14, caps=True, name=name)


# ───────────────────────── measuring (PIL) ─────────────────────────
_font_cache = {}


def pil_font(key, size):
    from PIL import ImageFont
    fam, _b, _w, var, _css, mono = FONTS[key]
    path = FONT_DIR / ('JetBrainsMono-Variable.ttf' if mono else 'Inter-Variable.ttf')
    k = (key, size)
    if k not in _font_cache:
        f = ImageFont.truetype(str(path), size)
        try:
            f.set_variation_by_name(var)
        except Exception:
            names = [n.decode() if isinstance(n, bytes) else n for n in f.get_variation_names()]
            if var in names:
                f.set_variation_by_name(var)
        _font_cache[k] = f
    return _font_cache[k]


def natural_lh(key, size):
    """Ascent+descent of the font at size, in px (what PowerPoint calls single spacing)."""
    a, d = pil_font(key, size).getmetrics()
    return a + d


def run_width(run):
    txt = run.text.upper() if run.caps else run.text
    f = pil_font(run.font, run.size)
    return f.getlength(txt) + run.spc * run.size * len(txt)


def wrap_para(para, width):
    """Greedy wrap over runs; returns list of lines, each a list of (Run, text) fragments."""
    lines, cur, cur_w = [], [], 0.0
    for run in para.runs:
        words = re.split(r'(\s+)', run.text)
        for wd in words:
            if wd == '':
                continue
            if '\n' in wd:
                wd = wd.replace('\n', ' ')
            frag = Run(wd, run.font, run.size, run.color, run.spc, run.caps, run.italic, run.url)
            fw = run_width(frag)
            if cur and cur_w + fw > width and wd.strip():
                lines.append(cur); cur, cur_w = [], 0.0
                if not wd.strip():
                    continue
            if not cur and not wd.strip():
                continue
            cur.append((frag, wd)); cur_w += fw
    if cur:
        lines.append(cur)
    return lines or [[]]


def text_height(el):
    total = 0.0
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        n = len(wrap_para(p, el.w))
        total += p.before + n * p.lh * size
    return total


# ───────────────────────── HTML backend ─────────────────────────
CSS = """
@font-face{font-family:'Inter';src:url('../vendor/fonts/Inter-Variable.ttf') format('truetype');font-weight:100 900;font-display:swap}
@font-face{font-family:'JetBrains Mono';src:url('../vendor/fonts/JetBrainsMono-Variable.ttf') format('truetype');font-weight:100 800;font-display:swap}
:root{--font-d:'Inter','Helvetica Now','Helvetica Neue',Helvetica,Arial,sans-serif;--font-m:'JetBrains Mono',Menlo,Consolas,monospace}
html,body{background:#000B1C}
.reveal{font-family:var(--font-d);color:#000B1C;-webkit-font-smoothing:antialiased;font-feature-settings:"ss01","cv11"}
.reveal .slides{text-align:left}
.reveal .slides section{padding:0;width:1920px;height:1080px;top:0;left:0;position:absolute;overflow:hidden}
.el{position:absolute;box-sizing:border-box;margin:0;display:flex;flex-direction:column;overflow:visible}
.el p{margin:0;white-space:pre-wrap;overflow-wrap:break-word}
.el p.bullet{padding-left:48px;text-indent:-48px}
.el p.bullet::before{content:'·';color:#ED6D24;font-family:var(--font-d);font-weight:900;display:inline-block;width:48px;text-indent:0}
.el a{color:inherit;text-decoration:underline;text-decoration-thickness:2px;text-underline-offset:.12em}
.el a:hover{color:#00544C}
.img{position:absolute;overflow:hidden}
.img img{width:100%;height:100%;display:block}
.fig{position:absolute}
.fig svg{width:100%;height:100%;display:block}
.embed{position:absolute;background:#000}
.embed iframe{width:100%;height:100%;border:0}
.cp{position:absolute;left:1450px;top:908px;width:350px;height:92px;border:2px dashed #ED6D24;color:#ED6D24;font:500 22px/1 var(--font-m);letter-spacing:.14em;text-transform:uppercase;display:flex;align-items:center;justify-content:center;gap:12px}
.cp b{width:12px;height:12px;border-radius:50%;background:#ED6D24;display:inline-block}
.reveal .progress{height:4px;color:#ED6D24}
.reveal .backgrounds{background:#000B1C}
"""

HTML_TMPL = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<link rel="stylesheet" href="../vendor/reveal/reset.css">
<link rel="stylesheet" href="../vendor/reveal/reveal.css">
<style>{css}</style>
</head>
<body>
<div class="reveal"><div class="slides">
{slides}
</div></div>
<script src="../vendor/reveal/reveal.js"></script>
<script src="../vendor/reveal/plugin/notes/notes.js"></script>
<script>
Reveal.initialize({{width:1920,height:1080,margin:0,minScale:0.05,maxScale:4,center:false,hash:true,transition:'none',
  backgroundTransition:'none',controls:false,progress:true,slideNumber:false,plugins:[RevealNotes],
  keyboard:{{}}}});
</script>
</body>
</html>
"""


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def css_font(run):
    fam, bold, weight, _v, css_fam, mono = FONTS[run.font]
    style = f"font-family:'{fam}',{'var(--font-m)' if mono else 'var(--font-d)'};font-weight:{weight};font-size:{run.size}px;color:{run.color};"
    if run.spc:
        style += f'letter-spacing:{run.spc}em;'
    if run.caps:
        style += 'text-transform:uppercase;'
    if run.italic:
        style += 'font-style:italic;'
    return style


def html_text(el):
    valign = {'t': 'flex-start', 'm': 'center', 'b': 'flex-end'}[el.valign]
    out = [f'<div class="el" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px;justify-content:{valign}">']
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        align = {'l': 'left', 'c': 'center', 'r': 'right'}[p.align]
        cls = ' class="bullet"' if p.bullet else ''
        out.append(f'<p{cls} style="text-align:{align};line-height:{p.lh};margin-top:{p.before}px;font-size:{size}px">')
        for r in p.runs:
            t = esc(r.text)
            span = f'<span style="{css_font(r)}">{t}</span>'
            if r.url:
                span = f'<a href="{esc(r.url)}" target="_blank" rel="noopener">{span}</a>'
            out.append(span)
        out.append('</p>')
    out.append('</div>')
    return ''.join(out)


def html_slide(s, i, assets_out, assets_rel):
    bg = f' data-background-color="{s.bg}"'
    parts = [f'<section{bg} data-slide="{i}">']
    for el in s.els + s.html_only:
        if el.kind == 'text':
            parts.append(html_text(el))
        elif el.kind == 'rect':
            st = f'left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px;'
            if el.fill:
                st += f'background:{el.fill};'
            if el.stroke:
                st += f'border:{el.stroke_w}px solid {el.stroke};'
            parts.append(f'<div class="el" style="{st}"></div>')
        elif el.kind == 'image':
            src = copy_asset(el.src, assets_out)
            fit = 'cover' if el.fit == 'cover' else 'contain'
            parts.append(f'<div class="img" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px"><img src="{assets_rel}/{src}" alt="" style="object-fit:{fit}"></div>')
        elif el.kind == 'figure':
            parts.append(f'<div class="fig" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px">{el.svg}</div>')
        elif el.kind == 'embed':
            parts.append(f'<div class="embed" style="left:{el.x}px;top:{el.y}px;width:{el.w}px;height:{el.h}px"><iframe data-src="https://www.youtube-nocookie.com/embed/{el.yt}?rel=0" allow="accelerometer; autoplay; clipboard-write; encrypted-media; gyroscope; picture-in-picture" allowfullscreen loading="lazy"></iframe></div>')
    if s.cp:
        label = {'word_cloud': 'Word cloud', 'multiple_choice': 'Multiple choice', 'short_answer': 'Short answer'}[s.cp['type']]
        parts.append(f'<div class="cp" data-classpoint><b></b>{label}</div>')
    if s.notes:
        parts.append(f'<aside class="notes">{esc(s.notes)}</aside>')
    parts.append('</section>')
    return '\n'.join(parts)


def copy_asset(src, assets_out):
    """Copy an image into the html assets dir (downscaled to <= 1920px, jpeg where possible)."""
    from PIL import Image as PImage
    p = Path(src) if Path(src).is_absolute() else ROOT / 'deck' / 'assets' / src
    assets_out.mkdir(parents=True, exist_ok=True)
    im = PImage.open(p)
    name = p.stem + ('.png' if im.mode in ('RGBA', 'LA', 'P') and p.suffix.lower() == '.png' and _has_alpha(im) else '.jpg')
    dst = assets_out / name
    if not dst.exists():
        im = im.convert('RGBA') if name.endswith('.png') else im.convert('RGB')
        if max(im.size) > 1920:
            im.thumbnail((1920, 1920))
        im.save(dst, quality=86, optimize=True) if name.endswith('.jpg') else im.save(dst, optimize=True)
    return name


def _has_alpha(im):
    return im.mode in ('RGBA', 'LA') and im.getextrema()[-1][0] < 255


def build_html(deck, out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    assets_out = out_dir / 'assets'
    slides = '\n'.join(html_slide(s, i + 1, assets_out, 'assets') for i, s in enumerate(deck['slides']))
    html = HTML_TMPL.format(title=esc(deck['title']), css=CSS, slides=slides)
    (out_dir / 'index.html').write_text(html, encoding='utf-8')
    vendor_fonts = out_dir.parent / 'vendor' / 'fonts'
    vendor_fonts.mkdir(parents=True, exist_ok=True)
    for f in FONT_DIR.glob('*.ttf'):
        shutil.copy(f, vendor_fonts / f.name)
    return out_dir / 'index.html'


# ───────────────────────── PPTX backend ─────────────────────────
def build_pptx(deck, out_path: Path, footer: str):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    master = prs.slide_masters[0]
    blank = next(lo for lo in master.slide_layouts if lo.name == 'Blank')
    # keep a single layout: build.py installs the ait4x master + 8 layouts next to it
    for lo in list(master.slide_layouts):
        if lo is blank:
            continue
        rId = master.part.relate_to(lo.part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout')
        for el in list(master.element.sldLayoutIdLst):
            if el.rId == rId:
                master.element.sldLayoutIdLst.remove(el)
        master.part.drop_rel(rId)

    def rgb(c):
        return RGBColor.from_string(c.lstrip('#'))

    def add_text(slide, el):
        tb = slide.shapes.add_textbox(Emu(el.x * PX), Emu(el.y * PX), Emu(el.w * PX), Emu(el.h * PX))
        if el.name:
            tb.name = el.name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[el.valign]
        for i, p in enumerate(el.paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[p.align]
            size = max((r.size for r in p.runs), default=24)
            key = p.runs[0].font if p.runs else 'body'
            # css line-height is a multiple of font size; PowerPoint's is a multiple of the font's natural height
            para.line_spacing = round(p.lh * size / natural_lh(key, size), 3)
            if p.before:
                para.space_before = Pt(p.before * 0.75)
            if p.bullet:
                pPr = para._p.get_or_add_pPr()
                pPr.set('marL', str(48 * PX)); pPr.set('indent', str(-48 * PX))
                buClr = pPr.makeelement(qn('a:buClr'), {}); clr = buClr.makeelement(qn('a:srgbClr'), {'val': 'ED6D24'}); buClr.append(clr)
                buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Inter Black'})
                buChar = pPr.makeelement(qn('a:buChar'), {'char': '·'})
                for e in (buClr, buFont, buChar):
                    pPr.append(e)
            for r in p.runs:
                run = para.add_run()
                run.text = r.text.upper() if r.caps else r.text
                fam, bold, _w, _v, _c, _m = FONTS[r.font]
                f = run.font
                f.name = fam
                f.size = Pt(r.size * 0.75)
                f.bold = bold
                f.italic = r.italic or None
                f.color.rgb = rgb(r.color)
                rPr = run._r.get_or_add_rPr()
                if r.spc:
                    rPr.set('spc', str(int(round(r.spc * r.size * 0.75 * 100))))
                if r.url:
                    run.hyperlink.address = r.url
        return tb

    def add_rect(slide, el):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(el.x * PX), Emu(el.y * PX), Emu(el.w * PX), Emu(el.h * PX))
        if el.name:
            sh.name = el.name
        if el.fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(el.fill)
        else:
            sh.fill.background()
        if el.stroke:
            sh.line.color.rgb = rgb(el.stroke); sh.line.width = Emu(el.stroke_w * PX)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def add_image(slide, path, el):
        from PIL import Image as PImage
        p = Path(path) if Path(path).is_absolute() else ROOT / 'deck' / 'assets' / path
        im = PImage.open(p)
        iw, ih = im.size
        box_ar, img_ar = el.w / el.h, iw / ih
        x, y, w, h = el.x, el.y, el.w, el.h
        if getattr(el, 'fit', 'cover') == 'contain':
            if img_ar > box_ar:
                h = round(w / img_ar); y = el.y + (el.h - h) // 2
            else:
                w = round(h * img_ar); x = el.x + (el.w - w) // 2
            pic = slide.shapes.add_picture(str(p), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        else:
            pic = slide.shapes.add_picture(str(p), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
            if img_ar > box_ar:      # too wide: crop left/right
                keep = box_ar / img_ar
                pic.crop_left = pic.crop_right = round((1 - keep) / 2, 4)
            elif img_ar < box_ar:    # too tall: crop top/bottom
                keep = img_ar / box_ar
                pic.crop_top = pic.crop_bottom = round((1 - keep) / 2, 4)
        if el.name:
            pic.name = el.name
        return pic

    for i, s in enumerate(deck['slides'], start=1):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = rgb(s.bg)
        for el in s.els:
            if el.kind == 'text':
                add_text(slide, el)
            elif el.kind == 'rect':
                add_rect(slide, el)
            elif el.kind == 'image':
                add_image(slide, el.src, el)
            elif el.kind == 'figure':
                add_image(slide, el.png, el)
            elif el.kind == 'embed':
                if el.thumb:
                    add_image(slide, el.thumb, el)
                url = f'https://www.youtube.com/watch?v={el.yt}'
                add_text(slide, Text(el.x, el.y + el.h + 12, el.w, 36, [Para(runs(f'[{url}]({url})', 'mono', 22, MUTED), 'l', 1.3)]))
        notes = s.notes or ''
        slide.notes_slide.notes_text_frame.text = notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    _normalise_layout(out_path)
    # ClassPoint manifest for build.py (1-based slide numbers)
    manifest = {'_comment': f'Generated from the deck spec — slide number -> ClassPoint activity. Footer: {footer}'}
    for i, s in enumerate(deck['slides'], start=1):
        if s.cp:
            manifest[str(i)] = s.cp
    manifest_path = out_path.with_name(out_path.stem + '-activities.json')
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding='utf-8')
    return out_path, manifest_path


def _normalise_layout(pptx_path: Path):
    """python-pptx keeps the template's layout number (slideLayout7); build.py expects slideLayout1."""
    with zipfile.ZipFile(pptx_path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    layouts = [n for n in parts if re.match(r'ppt/slideLayouts/slideLayout\d+\.xml$', n)]
    assert len(layouts) == 1, layouts
    old = re.search(r'slideLayout(\d+)\.xml', layouts[0]).group(1)
    if old != '1':
        def ren(n):
            return re.sub(rf'slideLayout{old}\.xml', 'slideLayout1.xml', n)
        parts = {ren(n): d for n, d in parts.items()}
        for n in list(parts):
            if n.endswith('.rels') or n == '[Content_Types].xml':
                parts[n] = parts[n].decode('utf-8').replace(f'slideLayout{old}.xml', 'slideLayout1.xml').encode('utf-8')
    with zipfile.ZipFile(pptx_path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)


def run_classpoint_build(pptx_path: Path, manifest_path: Path, footer: str):
    env = dict(os.environ, DECK_FOOTER=footer, DECK_ACTIVITIES=str(manifest_path))
    subprocess.run([sys.executable, str(ROOT / 'tools' / 'classpoint' / 'build.py'), str(pptx_path)], check=True, env=env)
    return pptx_path.with_name(pptx_path.stem + '-classpoint.pptx')


# ───────────────────────── PNG preview + overflow check ─────────────────────────
def build_png(deck, out_dir: Path, scale=0.5):
    from PIL import Image as PImage, ImageDraw
    out_dir.mkdir(parents=True, exist_ok=True)
    warnings = []
    files = []
    for i, s in enumerate(deck['slides'], start=1):
        im = PImage.new('RGB', (W, H), s.bg)
        d = ImageDraw.Draw(im)
        for el in s.els:
            if el.kind == 'rect':
                if el.fill:
                    d.rectangle([el.x, el.y, el.x + el.w - 1, el.y + el.h - 1], fill=el.fill)
                if el.stroke:
                    d.rectangle([el.x, el.y, el.x + el.w - 1, el.y + el.h - 1], outline=el.stroke, width=el.stroke_w)
            elif el.kind in ('image', 'figure', 'embed'):
                src = el.src if el.kind == 'image' else (el.png if el.kind == 'figure' else el.thumb)
                if not src:
                    d.rectangle([el.x, el.y, el.x + el.w, el.y + el.h], fill='#111')
                    continue
                p = Path(src) if Path(src).is_absolute() else ROOT / 'deck' / 'assets' / src
                pim = PImage.open(p).convert('RGBA')
                fit = getattr(el, 'fit', 'cover')
                pim = _fit(pim, el.w, el.h, fit)
                ox = el.x + (el.w - pim.width) // 2; oy = el.y + (el.h - pim.height) // 2
                im.paste(pim, (ox, oy), pim)
            elif el.kind == 'text':
                th = text_height(el)
                if th > el.h + 2:
                    warnings.append(f'slide {i:>2}: text overflows box by {th - el.h:.0f}px — "{el.paras[0].runs[0].text[:50] if el.paras and el.paras[0].runs else ""}"')
                _draw_text(d, el, th)
        if s.cp:
            d.rectangle([1450, 908, 1800, 1000], outline=ORANGE, width=3)
            f = pil_font('monomed', 22)
            d.text((1470, 940), 'CLASSPOINT · ' + s.cp['type'].replace('_', ' ').upper(), font=f, fill=ORANGE)
        if scale != 1:
            im = im.resize((int(W * scale), int(H * scale)), PImage.LANCZOS)
        fp = out_dir / f'slide-{i:02d}.png'
        im.save(fp)
        files.append(fp)
    return files, warnings


def _fit(pim, w, h, fit):
    from PIL import Image as PImage
    iw, ih = pim.size
    if fit == 'contain':
        r = min(w / iw, h / ih)
        return pim.resize((max(1, round(iw * r)), max(1, round(ih * r))), PImage.LANCZOS)
    r = max(w / iw, h / ih)
    pim = pim.resize((max(1, round(iw * r)), max(1, round(ih * r))), PImage.LANCZOS)
    l = (pim.width - w) // 2; t = (pim.height - h) // 2
    return pim.crop((l, t, l + w, t + h))


def _draw_text(d, el, th):
    y = el.y
    if el.valign == 'm':
        y = el.y + (el.h - th) / 2
    elif el.valign == 'b':
        y = el.y + el.h - th
    for p in el.paras:
        size = max((r.size for r in p.runs), default=24)
        lines = wrap_para(p, el.w)
        y += p.before
        for li, line in enumerate(lines):
            lw = sum(run_width(r) for r, _ in line)
            x = el.x + {'l': 0, 'c': (el.w - lw) / 2, 'r': el.w - lw}[p.align]
            if p.bullet:
                x += 48
                if li == 0:
                    d.text((el.x, y + (p.lh * size - size) / 2), '·', font=pil_font('black', size), fill=ORANGE)
            base = y + (p.lh * size - natural_lh(p.runs[0].font if p.runs else 'body', size)) / 2
            for r, _ in line:
                f = pil_font(r.font, r.size)
                txt = r.text.upper() if r.caps else r.text
                if r.spc:
                    cx = x
                    for ch in txt:
                        d.text((cx, base), ch, font=f, fill=r.color)
                        cx += f.getlength(ch) + r.spc * r.size
                    x = cx
                else:
                    d.text((x, base), txt, font=f, fill=r.color)
                    x += f.getlength(txt)
            y += p.lh * size


def contact_sheet(files, out, cols=4, scale=0.5):
    from PIL import Image as PImage, ImageDraw
    ims = [PImage.open(f) for f in files]
    if not ims:
        return
    w, h = ims[0].size
    rows = (len(ims) + cols - 1) // cols
    sheet = PImage.new('RGB', (cols * (w + 12) + 12, rows * (h + 34) + 12), '#DDD')
    d = ImageDraw.Draw(sheet)
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        x, y = 12 + c * (w + 12), 12 + r * (h + 34)
        sheet.paste(im, (x, y + 22))
        d.text((x, y + 4), f'{i + 1}', fill='#000', font=pil_font('monomed', 16))
    sheet.save(out, quality=80)
    return out


# ───────────────────────── orchestration ─────────────────────────
def build_all(deck, name: str, footer: str, do_pptx=True, do_html=True, do_png=True):
    outputs = {}
    if do_html:
        outputs['html'] = build_html(deck, ROOT / 'docs' / name)
    if do_pptx:
        pptx_path, manifest = build_pptx(deck, ROOT / 'export' / f'{name}.pptx', footer)
        outputs['pptx'] = pptx_path
        outputs['manifest'] = manifest
        outputs['classpoint'] = run_classpoint_build(pptx_path, manifest, footer)
    if do_png:
        files, warnings = build_png(deck, ROOT / 'export' / 'preview' / name)
        outputs['png'] = files
        outputs['warnings'] = warnings
        for k in range(0, len(files), 16):
            contact_sheet(files[k:k + 16], ROOT / 'export' / 'preview' / f'{name}-sheet-{k // 16 + 1}.jpg')
    return outputs
